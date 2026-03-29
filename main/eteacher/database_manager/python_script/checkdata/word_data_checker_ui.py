from __future__ import annotations

import importlib
import json
import re
import shutil
import sqlite3
import subprocess
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from difflib import SequenceMatcher, get_close_matches
from pathlib import Path
from typing import Any, Iterable


def _install_and_import(module_name: str, package_name: str | None = None) -> object:
    package = package_name or module_name
    try:
        return importlib.import_module(module_name)
    except ModuleNotFoundError:
        print(f"缺少依赖 {package}，正在尝试自动安装...", flush=True)
        result = subprocess.run(
            [sys.executable, "-m", "pip", "install", package],
            stdout=sys.stdout,
            stderr=sys.stderr,
            check=False,
        )
        if result.returncode != 0:
            raise RuntimeError(
                f"安装依赖 {package} 失败，请手动执行: {sys.executable} -m pip install {package}"
            ) from None
        return importlib.import_module(module_name)


def _try_import(module_name: str, package_name: str | None = None) -> object | None:
    try:
        return _install_and_import(module_name, package_name)
    except Exception:
        return None


xlrd = _install_and_import("xlrd")
xlwt = _install_and_import("xlwt")
openpyxl = _install_and_import("openpyxl")
wordfreq = _try_import("wordfreq")
eng_to_ipa = _try_import("eng_to_ipa")
_install_and_import("PySide6.QtCore", "PySide6")

from openpyxl import load_workbook
from PySide6.QtCore import QObject, QThread, Qt, Signal, Slot
from PySide6.QtGui import QColor
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QComboBox,
    QFileDialog,
    QGridLayout,
    QHeaderView,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QListWidget,
    QMainWindow,
    QMessageBox,
    QPlainTextEdit,
    QPushButton,
    QSpinBox,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QTreeWidget,
    QTreeWidgetItem,
    QVBoxLayout,
    QWidget,
)


DEFAULT_DB_PATH = Path(__file__).resolve().parents[1] / "words.db"
DEFAULT_REFERENCE_DB_PATH = DEFAULT_DB_PATH
DEFAULT_SEMANTIC_CACHE_PATH = Path(__file__).resolve().with_name("word_data_checker_semantic_cache.json")
EXACT_WORD_HEADERS = (
    "word",
    "phonetic",
    "word_type",
    "stage",
    "pos",
    "meaning_en",
    "meaning_zh",
    "image",
    "word_tag",
    "form_type",
    "form",
    "example_en",
    "example_zh",
    "difficulty",
    "example_tag",
    "audio_path",
    "image_path",
)

WHITESPACE_PATTERN = re.compile(r"\s+")
WORD_PATTERN = re.compile(r"[A-Za-z]+")
PHRASE_PATTERN = re.compile(r"[A-Za-z]+(?: [A-Za-z]+)+")
MEANING_ALLOWED_PATTERN = re.compile(
    r"^[A-Za-z0-9\u4E00-\u9FFF\u3400-\u4DBF\s"
    r"，。；：！？、,.!?;:/\\()（）\[\]【】{}<>《》\-+&%'\"“”‘’~·=…]+$"
)
PHONETIC_ALLOWED_PATTERN = re.compile(
    r"^[A-Za-z\u00C0-\u024F\u0250-\u02AF\u02B0-\u02FF\u0300-\u036F\u1D00-\u1D7F"
    r"\[\]/() '‘’`\-.,;:ːˈˌɚɝɡəæʌɑɔɒɛɜɪʊθðŋʃʒʧʤ]+$"
)
IPA_BODY_ALLOWED_PATTERN = re.compile(
    r"^[A-Za-z\u00C0-\u024F\u0250-\u02AF\u02B0-\u02FF\u0300-\u036F\u1D00-\u1D7F .,'ːˈˌ-]+$"
)
ASCII_ONLY_PATTERN = re.compile(r"^[A-Za-z' .:-]+$")
POS_PREFIX_PATTERN = re.compile(r"^\s*([A-Za-z.&]+)\s*")
MEANING_SPLIT_PATTERN = re.compile(r"[;；/、,，()（）\[\]【】{}<>《》\s]+")
WORD_ILLEGAL_CHARS_PATTERN = re.compile(r"[^A-Za-z ]")
PHONETIC_BRACKETS_PATTERN = re.compile(r"^(\[[^\]]+\]|/[^/]+/)$")

DEFAULT_TOO_MANY_MEANINGS = 8
SPELLING_WORDLIST_SIZE = 50000
IPA_EQUIVALENTS = [
    ("eɪ", "ei"),
    ("ɪ", "i"),
    ("ə", "ʌ"),
]

CATEGORY_TITLES = {
    "excel_header_invalid": "Excel 表头不合法",
    "word_empty": "word 字段为空",
    "word_invalid_word_chars": "单词包含非字母字符",
    "word_word_spacing": "单词空格不规范",
    "word_word_uppercase": "单词包含大写字母",
    "word_invalid_phrase_chars": "词组包含非法字符",
    "word_phrase_spacing": "词组空格数量不合理",
    "word_phrase_uppercase": "词组包含大写字母",
    "word_duplicate_word": "重复单词",
    "word_duplicate_phrase": "重复词组",
    "word_spelling_invalid": "单词拼写可疑",
    "phonetic_phrase_should_be_empty": "词组 phonetic 应为空",
    "phonetic_word_empty": "单词 phonetic 为空",
    "phonetic_invalid_chars": "phonetic 包含非音标字符",
    "phonetic_not_standard_ipa": "音标缺少 [] 或 // 包裹",
    "phonetic_word_mismatch": "音标与单词不一致",
    "meaning_zh_empty": "meaning_zh 为空",
    "meaning_zh_invalid_chars": "meaning_zh 包含非法字符",
    "meaning_semantic_mismatch": "中文释义与英文不匹配",
    "word_type_mismatch": "词性不一致",
    "word_without_meaning": "只有单词没有释义",
    "meaning_without_word": "只有释义没有单词",
    "too_many_meanings": "单词释义过多",
}

POS_ALIASES = {
    "n": "noun",
    "noun": "noun",
    "v": "verb",
    "vt": "verb",
    "vi": "verb",
    "verb": "verb",
    "adj": "adjective",
    "adjective": "adjective",
    "adv": "adverb",
    "adverb": "adverb",
    "prep": "preposition",
    "preposition": "preposition",
    "conj": "conjunction",
    "conjunction": "conjunction",
    "pron": "pronoun",
    "pronoun": "pronoun",
    "num": "numeral",
    "numeral": "numeral",
    "art": "article",
    "article": "article",
    "int": "interjection",
    "interj": "interjection",
    "interjection": "interjection",
    "aux": "auxiliary",
    "auxiliary": "auxiliary",
    "det": "determiner",
    "determiner": "determiner",
}

BUILTIN_MEANING_SEEDS = {
    "apple": {"苹果", "苹果树"},
    "teacher": {"教师", "老师"},
    "student": {"学生"},
    "car": {"汽车", "小汽车", "轿车"},
    "bus": {"公共汽车", "公交车", "总线"},
    "book": {"书", "书本", "预订", "登记"},
    "good": {"好的", "善良的", "有益的"},
    "bad": {"坏的", "糟糕的"},
    "teacher's": {"教师的"},
}

COLOR_BY_PREFIX = {
    "phonetic": QColor("#FFF4E5"),
    "meaning": QColor("#FDECEC"),
    "word": QColor("#EEF6FF"),
}


@dataclass(slots=True)
class WordRecord:
    source_path: Path
    source_name: str
    row_number: int | None
    record_id: int | None
    word_text: str
    phonetic: str
    word_type: str
    meaning_hint: str = ""
    pos_hint: str = ""


@dataclass(slots=True)
class MeaningRecord:
    source_path: Path
    source_name: str
    row_number: int | None
    record_id: int | None
    word_id: int | None
    word_text: str
    pos: str
    meaning_en: str
    meaning_zh: str
    word_type: str = ""


@dataclass(slots=True)
class FixAction:
    source_path: Path
    row_number: int
    field_name: str
    original_value: str
    new_value: str
    reason: str


@dataclass(slots=True)
class ValidationIssue:
    category_key: str
    category_title: str
    source_path: Path
    source_name: str
    table_name: str
    record_id: int | None
    word_id: int | None
    row_number: int | None
    word_text: str
    field_name: str
    raw_value: str
    normalized_value: str
    reason: str
    suggestion: str
    auto_fix_value: str = ""


@dataclass(slots=True)
class ValidationSummary:
    source_type: str
    source_label: str
    total_words: int
    total_meanings: int
    issues: list[ValidationIssue]
    fix_actions: list[FixAction]
    source_paths: list[Path]
    log_messages: list[str]


@dataclass(slots=True)
class ReferenceEntry:
    phonetics: set[str]
    pos_tokens: set[str]
    meaning_tokens: set[str]


def to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def normalize_spaces(value: object) -> str:
    return WHITESPACE_PATTERN.sub(" ", to_text(value).strip())


def is_blank(value: object) -> bool:
    return not normalize_spaces(value)


def contains_uppercase(text: str) -> bool:
    return any(character.isalpha() and character.isupper() for character in text)


def is_phrase(raw_word: str, normalized_word: str) -> bool:
    return " " in raw_word.strip() or " " in normalized_word


def find_invalid_characters(text: str, allowed_pattern: re.Pattern[str]) -> str:
    invalid_characters: list[str] = []
    for character in text:
        if not allowed_pattern.fullmatch(character):
            invalid_characters.append(character)

    unique_characters: list[str] = []
    for character in invalid_characters:
        if character not in unique_characters:
            unique_characters.append(character)
    return " ".join(repr(character) for character in unique_characters)


def canonical_pos_token(token: str) -> str:
    normalized = re.sub(r"[^A-Za-z]", "", token).lower()
    return POS_ALIASES.get(normalized, "")


def extract_pos_tokens(pos_text: str, meaning_zh: str = "") -> set[str]:
    tokens: set[str] = set()
    for part in re.split(r"[\s,，;；/、&]+", pos_text):
        canonical = canonical_pos_token(part)
        if canonical:
            tokens.add(canonical)

    prefix_match = POS_PREFIX_PATTERN.match(meaning_zh)
    if prefix_match is not None:
        canonical = canonical_pos_token(prefix_match.group(1))
        if canonical:
            tokens.add(canonical)
    return tokens


def strip_meaning_prefix(meaning_zh: str) -> str:
    return POS_PREFIX_PATTERN.sub("", meaning_zh, count=1).strip()


def extract_meaning_tokens(meaning_zh: str) -> set[str]:
    body = strip_meaning_prefix(meaning_zh)
    tokens: set[str] = set()
    for part in MEANING_SPLIT_PATTERN.split(body):
        token = part.strip(" .。;；:：,，!?！？")
        if not token:
            continue
        if not re.search(r"[\u4E00-\u9FFF]", token):
            continue
        if len(token) > 12:
            continue
        tokens.add(token)
    return tokens


def sanitize_word(raw_word: str, phrase: bool) -> str:
    normalized = normalize_spaces(raw_word)
    if not normalized:
        return ""
    cleaned = WORD_ILLEGAL_CHARS_PATTERN.sub(" " if phrase else "", normalized)
    cleaned = normalize_spaces(cleaned).lower()
    if phrase:
        return cleaned
    return cleaned.replace(" ", "")


def strip_invalid_meaning_characters(text: str) -> str:
    cleaned: list[str] = []
    for character in normalize_spaces(text):
        if MEANING_ALLOWED_PATTERN.fullmatch(character):
            cleaned.append(character)
    return normalize_spaces("".join(cleaned))


def strip_phonetic_brackets(phonetic: str) -> tuple[str, str, str]:
    normalized = normalize_spaces(phonetic)
    if normalized.startswith("[") and normalized.endswith("]"):
        return "[", normalized[1:-1].strip(), "]"
    if normalized.startswith("/") and normalized.endswith("/"):
        return "/", normalized[1:-1].strip(), "/"
    return "", normalized, ""


def normalize_ipa(ipa: str) -> str:
    return (
        ipa.replace("u:", "uː")
        .replace("i:", "iː")
        .replace("a:", "aː")
        .replace(":", "ː")
        .replace("ei", "eɪ")
        .replace("ai", "aɪ")
        .replace("oi", "ɔɪ")
        .replace("au", "aʊ")
        .replace("ou", "əʊ")
    )


def normalize_phonetic_body(body: str) -> str:
    normalized = normalize_spaces(body)
    replacements = {
        "（": "(",
        "）": ")",
        "【": "[",
        "】": "]",
        "：": ":",
        "'": "ˈ",
        "‘": "ˈ",
        "’": "ˈ",
        "g": "ɡ",
    }
    for old, new in replacements.items():
        normalized = normalized.replace(old, new)
    normalized = normalize_ipa(normalized)
    normalized = WHITESPACE_PATTERN.sub("", normalized)
    return normalized


def sanitize_phonetic(raw_phonetic: str) -> str:
    left, body, right = strip_phonetic_brackets(raw_phonetic)
    body = normalize_phonetic_body(body)
    filtered = "".join(character for character in body if IPA_BODY_ALLOWED_PATTERN.fullmatch(character))
    filtered = filtered.strip()
    if not filtered:
        return ""
    wrapper_left = left or "["
    wrapper_right = right or "]"
    return f"{wrapper_left}{filtered}{wrapper_right}"


def normalize_ipa_for_compare(phonetic: str) -> str:
    _, body, _ = strip_phonetic_brackets(phonetic)
    normalized = normalize_phonetic_body(body)
    normalized = normalized.replace("ɡ", "g")
    for canonical, equivalent in IPA_EQUIVALENTS:
        normalized = normalized.replace(equivalent, canonical)
    return normalized


def has_phonetic_wrapper(phonetic: str) -> bool:
    normalized = normalize_spaces(phonetic)
    if not normalized:
        return False

    return (
        len(normalized) >= 2
        and ((normalized.startswith("[") and normalized.endswith("]")) or (normalized.startswith("/") and normalized.endswith("/")))
    )


def ipa_similarity(left: str, right: str) -> float:
    return SequenceMatcher(None, normalize_ipa_for_compare(left), normalize_ipa_for_compare(right)).ratio()


class ReferenceLexicon:
    def __init__(self, reference_db_path: Path | None, cache_path: Path) -> None:
        self.reference_db_path = reference_db_path
        self.cache_path = cache_path
        self.entries: dict[str, ReferenceEntry] = {}
        self.common_words: list[str] = []
        self.common_word_set: set[str] = set()
        self._load_common_words()
        self._load_builtin_seeds()
        self._load_cache_file()
        self._load_reference_database()

    def _load_common_words(self) -> None:
        if wordfreq is None:
            return
        try:
            common_words = list(wordfreq.top_n_list("en", SPELLING_WORDLIST_SIZE))
        except Exception:
            return
        self.common_words = [normalize_spaces(word).lower() for word in common_words if normalize_spaces(word)]
        self.common_word_set = set(self.common_words)

    def _ensure_entry(self, word: str) -> ReferenceEntry:
        key = word.casefold()
        entry = self.entries.get(key)
        if entry is None:
            entry = ReferenceEntry(set(), set(), set())
            self.entries[key] = entry
        return entry

    def _load_builtin_seeds(self) -> None:
        for word, meanings in BUILTIN_MEANING_SEEDS.items():
            entry = self._ensure_entry(word)
            entry.meaning_tokens.update(meanings)

    def _load_cache_file(self) -> None:
        if not self.cache_path.exists():
            return
        try:
            payload = json.loads(self.cache_path.read_text(encoding="utf-8"))
        except Exception:
            return
        if not isinstance(payload, dict):
            return
        for word, value in payload.items():
            if not isinstance(value, dict):
                continue
            entry = self._ensure_entry(word)
            for phonetic in value.get("phonetics", []):
                normalized = normalize_spaces(phonetic)
                if normalized:
                    entry.phonetics.add(normalized)
            for pos in value.get("pos", []):
                entry.pos_tokens.update(extract_pos_tokens(normalize_spaces(pos)))
            for meaning in value.get("meanings_zh", []):
                normalized_meaning = normalize_spaces(meaning)
                if normalized_meaning:
                    entry.meaning_tokens.update(extract_meaning_tokens(normalized_meaning))

    def _load_reference_database(self) -> None:
        if self.reference_db_path is None or not self.reference_db_path.exists():
            return
        with sqlite3.connect(self.reference_db_path) as conn:
            conn.row_factory = sqlite3.Row
            rows = conn.execute(
                """
                SELECT
                    COALESCE(w.word, '') AS word,
                    COALESCE(w.phonetic, '') AS phonetic,
                    COALESCE(w.word_type, '') AS word_type,
                    COALESCE(wm.pos, '') AS pos,
                    COALESCE(wm.meaning_zh, '') AS meaning_zh
                FROM word AS w
                LEFT JOIN word_meaning AS wm ON wm.word_id = w.id
                ORDER BY w.word, wm.id
                """
            ).fetchall()

        for row in rows:
            word = normalize_spaces(row["word"]).lower()
            if not word:
                continue
            entry = self._ensure_entry(word)
            phonetic = normalize_spaces(row["phonetic"])
            if phonetic:
                entry.phonetics.add(phonetic)
            entry.pos_tokens.update(extract_pos_tokens(normalize_spaces(row["word_type"])))
            entry.pos_tokens.update(extract_pos_tokens(normalize_spaces(row["pos"]), normalize_spaces(row["meaning_zh"])))
            entry.meaning_tokens.update(extract_meaning_tokens(normalize_spaces(row["meaning_zh"])))

    def get(self, word: str) -> ReferenceEntry | None:
        return self.entries.get(word.casefold())

    def is_valid_word(self, word: str) -> bool:
        normalized = normalize_spaces(word).lower()
        if not normalized:
            return False
        if self.get(normalized) is not None:
            return True
        if wordfreq is not None:
            try:
                if float(wordfreq.zipf_frequency(normalized, "en")) >= 2.2:
                    return True
            except Exception:
                pass
        return normalized in self.common_word_set

    def suggest_word(self, word: str) -> str:
        normalized = normalize_spaces(word).lower()
        candidates: list[str] = []
        if self.common_words:
            candidates = get_close_matches(normalized, self.common_words, n=1, cutoff=0.86)
        if not candidates and self.entries:
            candidates = get_close_matches(normalized, list(self.entries.keys()), n=1, cutoff=0.86)
        return candidates[0] if candidates else ""

    def get_expected_phonetic(self, word: str) -> str:
        normalized = normalize_spaces(word).lower()
        entry = self.get(normalized)
        if entry is not None and entry.phonetics:
            return sorted(entry.phonetics, key=len)[0]
        if eng_to_ipa is not None:
            try:
                generated = normalize_spaces(eng_to_ipa.convert(normalized))
            except Exception:
                generated = ""
            if generated and "*" not in generated:
                return f"[{normalize_phonetic_body(generated)}]"
        return ""

    def semantic_matches(self, word: str, meaning_zh: str) -> bool | None:
        normalized = normalize_spaces(word).lower()
        entry = self.get(normalized)
        if entry is None or not entry.meaning_tokens:
            return None
        actual_tokens = extract_meaning_tokens(meaning_zh)
        if not actual_tokens:
            return None
        if actual_tokens & entry.meaning_tokens:
            return True
        for actual in actual_tokens:
            for expected in entry.meaning_tokens:
                if actual in expected or expected in actual:
                    return True
        return False


class ValidationService:
    def __init__(self, reference_db_path: Path | None, cache_path: Path, too_many_meanings: int) -> None:
        self.reference_lexicon = ReferenceLexicon(reference_db_path, cache_path)
        self.too_many_meanings = too_many_meanings

    def validate_database(self, db_path: Path) -> ValidationSummary:
        if not db_path.exists():
            raise FileNotFoundError(f"数据库不存在: {db_path}")

        with sqlite3.connect(db_path) as conn:
            conn.row_factory = sqlite3.Row
            word_rows = conn.execute(
                """
                SELECT id, COALESCE(word, '') AS word, COALESCE(phonetic, '') AS phonetic, COALESCE(word_type, '') AS word_type
                FROM word
                ORDER BY id
                """
            ).fetchall()
            meaning_rows = conn.execute(
                """
                SELECT
                    wm.id AS id,
                    wm.word_id AS word_id,
                    COALESCE(w.word, '') AS word,
                    COALESCE(w.word_type, '') AS word_type,
                    COALESCE(wm.pos, '') AS pos,
                    COALESCE(wm.meaning_en, '') AS meaning_en,
                    COALESCE(wm.meaning_zh, '') AS meaning_zh
                FROM word_meaning AS wm
                LEFT JOIN word AS w ON w.id = wm.word_id
                ORDER BY wm.id
                """
            ).fetchall()

        word_records = [
            WordRecord(
                source_path=db_path,
                source_name=db_path.name,
                row_number=None,
                record_id=int(row["id"]),
                word_text=to_text(row["word"]),
                phonetic=to_text(row["phonetic"]),
                word_type=to_text(row["word_type"]),
            )
            for row in word_rows
        ]
        meaning_records = [
            MeaningRecord(
                source_path=db_path,
                source_name=db_path.name,
                row_number=None,
                record_id=int(row["id"]),
                word_id=int(row["word_id"]) if row["word_id"] is not None else None,
                word_text=to_text(row["word"]),
                pos=to_text(row["pos"]),
                meaning_en=to_text(row["meaning_en"]),
                meaning_zh=to_text(row["meaning_zh"]),
                word_type=to_text(row["word_type"]),
            )
            for row in meaning_rows
        ]
        issues, logs = self._validate_records(word_records, meaning_records, source_type="db")
        return ValidationSummary(
            source_type="db",
            source_label=str(db_path),
            total_words=len(word_records),
            total_meanings=len(meaning_records),
            issues=issues,
            fix_actions=[],
            source_paths=[db_path],
            log_messages=logs,
        )

    def validate_excel_files(self, excel_paths: Iterable[Path]) -> ValidationSummary:
        valid_files = [path for path in excel_paths if path.exists() and not path.name.startswith("~$")]
        if not valid_files:
            raise ValueError("未选择有效的 Excel 文件。")

        word_records: list[WordRecord] = []
        meaning_records: list[MeaningRecord] = []
        issues: list[ValidationIssue] = []
        logs: list[str] = []

        for path in valid_files:
            file_issues, file_logs, file_word_records, file_meaning_records = self._read_excel_file(path)
            issues.extend(file_issues)
            logs.extend(file_logs)
            word_records.extend(file_word_records)
            meaning_records.extend(file_meaning_records)

        record_issues, record_logs = self._validate_records(word_records, meaning_records, source_type="excel")
        issues.extend(record_issues)
        logs.extend(record_logs)
        fix_actions = self._build_fix_actions(record_issues)

        return ValidationSummary(
            source_type="excel",
            source_label=f"{len(valid_files)} 个 Excel 文件",
            total_words=len(word_records),
            total_meanings=len(meaning_records),
            issues=issues,
            fix_actions=fix_actions,
            source_paths=valid_files,
            log_messages=logs,
        )

    def apply_fixed_excel_files(self, fix_actions: list[FixAction]) -> list[Path]:
        if not fix_actions:
            raise ValueError("当前没有可写回的自动修复项。")

        fixes_by_file: dict[Path, dict[tuple[int, str], FixAction]] = defaultdict(dict)
        for action in fix_actions:
            fixes_by_file[action.source_path][(action.row_number, action.field_name)] = action

        updated_files: list[Path] = []
        for source_path, indexed_actions in fixes_by_file.items():
            backup_path = self._build_backup_path(source_path)
            shutil.copy2(source_path, backup_path)

            if source_path.suffix.lower() == ".xlsx":
                self._apply_fixes_to_xlsx(source_path, indexed_actions)
            elif source_path.suffix.lower() == ".xls":
                self._apply_fixes_to_xls(source_path, indexed_actions)
            else:
                raise ValueError(f"不支持的 Excel 类型：{source_path.name}")

            updated_files.append(source_path)

        return updated_files

    def _build_backup_path(self, source_path: Path) -> Path:
        return source_path.parent / f"{source_path.stem}.bak{source_path.suffix}"

    def _build_fix_column_mapping(self, source_path: Path) -> tuple[dict[str, int], list[list[object]]]:
        rows = self._load_excel_rows(source_path)
        header_index, mapping, _, _ = self._detect_header(rows)
        if header_index is None:
            raise ValueError(f"{source_path.name} 表头识别失败，无法回写修复。")
        return mapping, rows

    def _apply_fixes_to_xlsx(
        self,
        source_path: Path,
        indexed_actions: dict[tuple[int, str], FixAction],
    ) -> None:
        mapping, _ = self._build_fix_column_mapping(source_path)
        workbook = load_workbook(source_path)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            for (row_number, field_name), action in indexed_actions.items():
                column_index = mapping.get(field_name)
                if column_index is None:
                    continue
                sheet.cell(row=row_number, column=column_index + 1, value=action.new_value)
            workbook.save(source_path)
        finally:
            workbook.close()

    def _apply_fixes_to_xls(
        self,
        source_path: Path,
        indexed_actions: dict[tuple[int, str], FixAction],
    ) -> None:
        mapping, rows = self._build_fix_column_mapping(source_path)
        mutable_rows = [list(row) for row in rows]
        for (row_number, field_name), action in indexed_actions.items():
            column_index = mapping.get(field_name)
            if column_index is None:
                continue
            row_index = row_number - 1
            if row_index < 0 or row_index >= len(mutable_rows):
                continue
            while len(mutable_rows[row_index]) <= column_index:
                mutable_rows[row_index].append("")
            mutable_rows[row_index][column_index] = action.new_value

        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet((source_path.stem or "Sheet1")[:31])
        for row_index, row in enumerate(mutable_rows):
            for column_index, value in enumerate(row):
                sheet.write(row_index, column_index, to_text(value))
        workbook.save(str(source_path))

    def _read_excel_file(
        self,
        path: Path,
    ) -> tuple[list[ValidationIssue], list[str], list[WordRecord], list[MeaningRecord]]:
        rows = self._load_excel_rows(path)
        if not rows:
            issue = ValidationIssue(
                category_key="excel_header_invalid",
                category_title=CATEGORY_TITLES["excel_header_invalid"],
                source_path=path,
                source_name=path.name,
                table_name="excel",
                record_id=None,
                word_id=None,
                row_number=None,
                word_text="",
                field_name="header",
                raw_value="",
                normalized_value="",
                reason="Excel 文件为空，未找到任何表头。",
                suggestion=f"表头必须存在，且字段名需与脚本字段名完全一致：{', '.join(EXACT_WORD_HEADERS)}",
            )
            return [issue], [f"{path.name}: Excel 文件为空。"], [], []

        header_index, mapping, unmatched_headers, unmatched_header_index = self._detect_header(rows)
        if header_index is None:
            reason = (
                f"前 5 行内未找到合法表头。"
                f"表头必须存在，且字段名需与脚本字段名完全一致：{', '.join(EXACT_WORD_HEADERS)}"
            )
            issue = ValidationIssue(
                category_key="excel_header_invalid",
                category_title=CATEGORY_TITLES["excel_header_invalid"],
                source_path=path,
                source_name=path.name,
                table_name="excel",
                record_id=None,
                word_id=None,
                row_number=unmatched_header_index + 1 if unmatched_header_index is not None else None,
                word_text="",
                field_name="header",
                raw_value=", ".join(unmatched_headers),
                normalized_value=", ".join(unmatched_headers),
                reason=reason,
                suggestion="请将 Excel 表头改成脚本支持的标准字段名。",
            )
            return [issue], [f"{path.name}: 表头识别失败。"], [], []

        logs = [f"{path.name}: 识别到表头，第 {header_index + 1} 行。"]
        if unmatched_headers:
            logs.append(f"{path.name}: 未匹配表头 -> {', '.join(unmatched_headers)}")

        word_records: list[WordRecord] = []
        meaning_records: list[MeaningRecord] = []
        for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if all(is_blank(cell) for cell in row):
                continue
            record = self._record_from_excel_row(path, row_number, row, mapping)
            if record is None:
                continue
            word_records.append(
                WordRecord(
                    source_path=path,
                    source_name=path.name,
                    row_number=row_number,
                    record_id=row_number,
                    word_text=to_text(record.get("word")),
                    phonetic=to_text(record.get("phonetic")),
                    word_type=to_text(record.get("word_type")),
                    meaning_hint=to_text(record.get("meaning_zh")),
                    pos_hint=to_text(record.get("pos")),
                )
            )
            meaning_records.append(
                MeaningRecord(
                    source_path=path,
                    source_name=path.name,
                    row_number=row_number,
                    record_id=row_number,
                    word_id=None,
                    word_text=to_text(record.get("word")),
                    pos=to_text(record.get("pos")),
                    meaning_en=to_text(record.get("meaning_en")),
                    meaning_zh=to_text(record.get("meaning_zh")),
                    word_type=to_text(record.get("word_type")),
                )
            )
        return [], logs, word_records, meaning_records

    def _load_excel_rows(self, path: Path) -> list[list[object]]:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = workbook[workbook.sheetnames[0]]
                return [list(row) for row in sheet.iter_rows(values_only=True)]
            finally:
                workbook.close()
        if suffix == ".xls":
            workbook = xlrd.open_workbook(path.as_posix())
            sheet = workbook.sheet_by_index(0)
            return [sheet.row_values(index) for index in range(sheet.nrows)]
        raise ValueError(f"不支持的 Excel 类型：{path.name}")

    def _detect_header(
        self,
        rows: list[list[object]],
    ) -> tuple[int | None, dict[str, int], list[str], int | None]:
        max_scan = min(5, len(rows))
        best_unmatched_headers: list[str] = []
        best_unmatched_index: int | None = None
        best_matched_count = -1
        for index in range(max_scan):
            mapping, unmatched_headers = self._build_column_map(rows[index])
            if len(mapping) > best_matched_count:
                best_matched_count = len(mapping)
                best_unmatched_headers = unmatched_headers
                best_unmatched_index = index if unmatched_headers else None
            elif best_unmatched_index is None and unmatched_headers:
                best_unmatched_headers = unmatched_headers
                best_unmatched_index = index
            if len(mapping) >= 2 and "word" in mapping:
                return index, mapping, unmatched_headers, index
        return None, {}, best_unmatched_headers, best_unmatched_index

    def _build_column_map(self, header_row: list[object]) -> tuple[dict[str, int], list[str]]:
        mapping: dict[str, int] = {}
        unmatched_headers: list[str] = []
        for index, value in enumerate(header_row):
            header_name = normalize_spaces(value)
            if not header_name:
                continue
            if header_name not in EXACT_WORD_HEADERS:
                if header_name not in unmatched_headers:
                    unmatched_headers.append(header_name)
                continue
            mapping.setdefault(header_name, index)
        return mapping, unmatched_headers

    def _record_from_excel_row(
        self,
        path: Path,
        row_number: int,
        row: list[object],
        mapping: dict[str, int],
    ) -> dict[str, object] | None:
        record: dict[str, object] = {
            "_file": path.name,
            "_row_number": row_number,
        }
        for name, column_index in mapping.items():
            record[name] = row[column_index] if column_index < len(row) else None
        if all(
            is_blank(record.get(key))
            for key in ("word", "meaning_zh", "meaning_en", "form", "example_en", "example_zh")
        ):
            return None
        return record

    def _validate_records(
        self,
        word_records: list[WordRecord],
        meaning_records: list[MeaningRecord],
        source_type: str,
    ) -> tuple[list[ValidationIssue], list[str]]:
        issues: list[ValidationIssue] = []
        logs: list[str] = []

        if source_type == "db":
            duplicate_words: dict[str, list[WordRecord]] = defaultdict(list)
            duplicate_phrases: dict[str, list[WordRecord]] = defaultdict(list)
        else:
            duplicate_words = defaultdict(list)
            duplicate_phrases = defaultdict(list)

        meaning_groups: dict[str, set[str]] = defaultdict(set)
        word_has_meaning: Counter[str] = Counter()

        for meaning_record in meaning_records:
            normalized_word = normalize_spaces(meaning_record.word_text).lower()
            normalized_meaning = normalize_spaces(meaning_record.meaning_zh)
            normalized_pos = normalize_spaces(meaning_record.pos)
            if normalized_word and normalized_meaning:
                meaning_groups[normalized_word].add(f"{normalized_pos}::{normalized_meaning}")
                word_has_meaning[normalized_word] += 1

        for record in word_records:
            issues.extend(self._validate_word_record(record, source_type, duplicate_words, duplicate_phrases, word_has_meaning))

        issues.extend(self._emit_duplicate_issues(duplicate_words, duplicate_phrases, source_type))

        for record in meaning_records:
            issues.extend(self._validate_meaning_record(record))

        for word_key, meanings in meaning_groups.items():
            if len(meanings) <= self.too_many_meanings:
                continue
            word_text = next((record.word_text for record in word_records if normalize_spaces(record.word_text).lower() == word_key), word_key)
            source_path = next((record.source_path for record in word_records if normalize_spaces(record.word_text).lower() == word_key), Path(word_text))
            source_name = next((record.source_name for record in word_records if normalize_spaces(record.word_text).lower() == word_key), "")
            row_number = next((record.row_number for record in word_records if normalize_spaces(record.word_text).lower() == word_key), None)
            issues.append(
                ValidationIssue(
                    category_key="too_many_meanings",
                    category_title=CATEGORY_TITLES["too_many_meanings"],
                    source_path=source_path,
                    source_name=source_name,
                    table_name="word_meaning",
                    record_id=None,
                    word_id=None,
                    row_number=row_number,
                    word_text=word_text,
                    field_name="meaning_zh",
                    raw_value=str(len(meanings)),
                    normalized_value=str(len(meanings)),
                    reason=f"单词 {word_text} 关联了 {len(meanings)} 个不同释义，超过阈值 {self.too_many_meanings}。",
                    suggestion="检查是否存在重复导入、脏数据或误拆分释义。",
                )
            )

        logs.append(f"校验完成，共发现 {len(issues)} 条问题。")
        return issues, logs

    def _validate_word_record(
        self,
        record: WordRecord,
        source_type: str,
        duplicate_words: dict[str, list[WordRecord]],
        duplicate_phrases: dict[str, list[WordRecord]],
        word_has_meaning: Counter[str],
    ) -> list[ValidationIssue]:
        issues: list[ValidationIssue] = []
        raw_word = to_text(record.word_text)
        raw_phonetic = to_text(record.phonetic)
        normalized_word = normalize_spaces(raw_word)
        normalized_phonetic = normalize_spaces(raw_phonetic)
        phrase = is_phrase(raw_word, normalized_word)

        if not normalized_word:
            if normalize_spaces(record.meaning_hint):
                issues.append(
                    self._issue(
                        category_key="meaning_without_word",
                        record=record,
                        table_name="excel" if record.row_number is not None else "word_meaning",
                        field_name="word",
                        raw_value=raw_word,
                        normalized_value=normalized_word,
                        reason="存在中文释义，但 word 字段为空。",
                        suggestion="补充对应英文单词，或删除孤立释义。",
                    )
                )
            else:
                issues.append(
                    self._issue(
                        category_key="word_empty",
                        record=record,
                        table_name="word",
                        field_name="word",
                        raw_value=raw_word,
                        normalized_value=normalized_word,
                        reason="word 字段去除空白后为空。",
                        suggestion="填写有效英文单词或词组。",
                    )
                )
            return issues

        lowered_key = normalized_word.casefold()
        if phrase:
            signature = self._duplicate_signature(record, phrase=True, source_type=source_type)
            duplicate_phrases[signature].append(record)
            if raw_word != normalized_word or "  " in raw_word:
                fixed_value = sanitize_word(raw_word, phrase=True)
                issues.append(
                    self._issue(
                        category_key="word_phrase_spacing",
                        record=record,
                        table_name="word",
                        field_name="word",
                        raw_value=raw_word,
                        normalized_value=normalized_word,
                        reason="词组存在首尾空格或连续空格。",
                        suggestion=f"建议改为: {fixed_value or normalized_word}",
                        auto_fix_value=fixed_value,
                    )
                )
            if contains_uppercase(raw_word):
                fixed_value = sanitize_word(raw_word, phrase=True)
                issues.append(
                    self._issue(
                        category_key="word_phrase_uppercase",
                        record=record,
                        table_name="word",
                        field_name="word",
                        raw_value=raw_word,
                        normalized_value=normalize_spaces(raw_word).lower(),
                        reason="词组中包含大写英文字母。",
                        suggestion=f"建议统一为小写: {fixed_value or normalized_word.lower()}",
                        auto_fix_value=fixed_value,
                    )
                )
            if not PHRASE_PATTERN.fullmatch(normalized_word):
                invalid_chars = find_invalid_characters(normalized_word, re.compile(r"[A-Za-z ]"))
                fixed_value = sanitize_word(raw_word, phrase=True)
                issues.append(
                    self._issue(
                        category_key="word_invalid_phrase_chars",
                        record=record,
                        table_name="word",
                        field_name="word",
                        raw_value=raw_word,
                        normalized_value=normalized_word,
                        reason=f"词组只能由英文单词和单个空格组成，检测到非法字符: {invalid_chars or '未知'}。",
                        suggestion="删除非英文字母字符，并确保单词之间只保留一个空格。",
                        auto_fix_value=fixed_value,
                    )
                )
        else:
            signature = self._duplicate_signature(record, phrase=False, source_type=source_type)
            duplicate_words[signature].append(record)
            if raw_word != normalized_word or " " in raw_word:
                fixed_value = sanitize_word(raw_word, phrase=False)
                issues.append(
                    self._issue(
                        category_key="word_word_spacing",
                        record=record,
                        table_name="word",
                        field_name="word",
                        raw_value=raw_word,
                        normalized_value=normalized_word,
                        reason="单词不应包含首尾空格或内部空格。",
                        suggestion=f"建议改为: {fixed_value or normalized_word.replace(' ', '')}",
                        auto_fix_value=fixed_value,
                    )
                )
            if contains_uppercase(raw_word):
                fixed_value = sanitize_word(raw_word, phrase=False)
                issues.append(
                    self._issue(
                        category_key="word_word_uppercase",
                        record=record,
                        table_name="word",
                        field_name="word",
                        raw_value=raw_word,
                        normalized_value=normalized_word.lower(),
                        reason="单词中包含大写英文字母。",
                        suggestion=f"建议统一为小写: {fixed_value or normalized_word.lower()}",
                        auto_fix_value=fixed_value,
                    )
                )
            if not WORD_PATTERN.fullmatch(normalized_word):
                invalid_chars = find_invalid_characters(normalized_word, re.compile(r"[A-Za-z]"))
                fixed_value = sanitize_word(raw_word, phrase=False)
                issues.append(
                    self._issue(
                        category_key="word_invalid_word_chars",
                        record=record,
                        table_name="word",
                        field_name="word",
                        raw_value=raw_word,
                        normalized_value=normalized_word,
                        reason=f"单词只能包含英文字母，检测到非法字符: {invalid_chars or '未知'}。",
                        suggestion="删除非字母字符。",
                        auto_fix_value=fixed_value,
                    )
                )

            clean_word = sanitize_word(raw_word, phrase=False)
            if clean_word and WORD_PATTERN.fullmatch(clean_word) and not self.reference_lexicon.is_valid_word(clean_word):
                suggestion = self.reference_lexicon.suggest_word(clean_word)
                issues.append(
                    self._issue(
                        category_key="word_spelling_invalid",
                        record=record,
                        table_name="word",
                        field_name="word",
                        raw_value=raw_word,
                        normalized_value=clean_word,
                        reason="该单词在参考词典和常用词表中都未命中，疑似拼写错误。",
                        suggestion=(f"可疑拼写，候选词: {suggestion}" if suggestion else "请核对单词拼写。"),
                    )
                )

        if phrase:
            if normalized_phonetic:
                issues.append(
                    self._issue(
                        category_key="phonetic_phrase_should_be_empty",
                        record=record,
                        table_name="word",
                        field_name="phonetic",
                        raw_value=raw_phonetic,
                        normalized_value=normalized_phonetic,
                        reason="词组的 phonetic 按当前规则应为空。",
                        suggestion="清空该词组的 phonetic 字段。",
                        auto_fix_value="",
                    )
                )
        else:
            if normalized_word and not normalized_phonetic:
                issues.append(
                    self._issue(
                        category_key="phonetic_word_empty",
                        record=record,
                        table_name="word",
                        field_name="phonetic",
                        raw_value=raw_phonetic,
                        normalized_value=normalized_phonetic,
                        reason="单词缺少 phonetic。",
                        suggestion="补充该单词的音标。",
                    )
                )
            elif normalized_phonetic:
                if not has_phonetic_wrapper(normalized_phonetic):
                    issues.append(
                        self._issue(
                            category_key="phonetic_not_standard_ipa",
                            record=record,
                            table_name="word",
                            field_name="phonetic",
                            raw_value=raw_phonetic,
                            normalized_value=sanitize_phonetic(raw_phonetic),
                            reason="音标只校验首尾是否使用 [] 或 // 包裹，当前格式不符合要求。",
                            suggestion="将音标改为使用 [] 或 // 包裹，例如 [abc] 或 /abc/。",
                            auto_fix_value=sanitize_phonetic(raw_phonetic),
                        )
                    )

        if normalized_word and not word_has_meaning[lowered_key] and record.row_number is not None:
            issues.append(
                self._issue(
                    category_key="word_without_meaning",
                    record=record,
                    table_name="excel",
                    field_name="meaning_zh",
                    raw_value=record.meaning_hint,
                    normalized_value=normalize_spaces(record.meaning_hint),
                    reason="当前行包含单词，但缺少有效中文释义。",
                    suggestion="补充 meaning_zh，或删除孤立单词行。",
                )
            )

        return issues

    def _validate_meaning_record(self, record: MeaningRecord) -> list[ValidationIssue]:
        issues: list[ValidationIssue] = []
        raw_meaning = to_text(record.meaning_zh)
        raw_word = to_text(record.word_text)
        normalized_word = normalize_spaces(raw_word)
        normalized_meaning = normalize_spaces(raw_meaning)

        if not normalized_word and normalized_meaning:
            issues.append(
                self._issue_from_meaning(
                    category_key="meaning_without_word",
                    record=record,
                    field_name="word",
                    raw_value=raw_word,
                    normalized_value=normalized_word,
                    reason="存在释义内容，但没有对应单词。",
                    suggestion="补充 word 字段，或删除孤立释义。",
                )
            )
            return issues

        if normalized_word and not normalized_meaning:
            issues.append(
                self._issue_from_meaning(
                    category_key="meaning_zh_empty",
                    record=record,
                    field_name="meaning_zh",
                    raw_value=raw_meaning,
                    normalized_value=normalized_meaning,
                    reason="meaning_zh 去除空白后为空。",
                    suggestion="补充有效中文释义。",
                )
            )
            return issues

        if not MEANING_ALLOWED_PATTERN.fullmatch(normalized_meaning):
            invalid_chars = find_invalid_characters(normalized_meaning, MEANING_ALLOWED_PATTERN)
            fixed_value = strip_invalid_meaning_characters(raw_meaning)
            issues.append(
                self._issue_from_meaning(
                    category_key="meaning_zh_invalid_chars",
                    record=record,
                    field_name="meaning_zh",
                    raw_value=raw_meaning,
                    normalized_value=normalized_meaning,
                    reason=f"meaning_zh 包含非法字符: {invalid_chars or '未知'}。",
                    suggestion="删除异常字符，保留中文释义、词性缩写和常见标点。",
                    auto_fix_value=fixed_value,
                )
            )

        word_type_tokens = extract_pos_tokens(normalize_spaces(record.word_type))
        meaning_pos_tokens = extract_pos_tokens(normalize_spaces(record.pos), normalized_meaning)
        if word_type_tokens and meaning_pos_tokens and not (word_type_tokens & meaning_pos_tokens):
            issues.append(
                self._issue_from_meaning(
                    category_key="word_type_mismatch",
                    record=record,
                    field_name="word_type",
                    raw_value=record.word_type,
                    normalized_value=record.word_type,
                    reason=(
                        f"word_type={', '.join(sorted(word_type_tokens))} 与释义词性="
                        f"{', '.join(sorted(meaning_pos_tokens))} 不一致。"
                    ),
                    suggestion="统一 word_type、pos 和 meaning_zh 前缀中的词性信息。",
                )
            )

        semantic_result = self.reference_lexicon.semantic_matches(normalized_word, normalized_meaning)
        if semantic_result is False:
            issues.append(
                self._issue_from_meaning(
                    category_key="meaning_semantic_mismatch",
                    record=record,
                    field_name="meaning_zh",
                    raw_value=raw_meaning,
                    normalized_value=normalized_meaning,
                    reason="当前中文释义与参考词典中的语义不一致。",
                    suggestion="核对该英文词的中文释义，确认是否录错或串行。",
                )
            )

        return issues

    def _duplicate_signature(self, record: WordRecord, phrase: bool, source_type: str) -> str:
        normalized_word = normalize_spaces(record.word_text).casefold()
        if source_type == "db":
            return normalized_word
        meaning_hint = normalize_spaces(record.meaning_hint).casefold()
        phonetic = normalize_spaces(record.phonetic).casefold()
        word_type = normalize_spaces(record.word_type).casefold()
        pos_hint = normalize_spaces(record.pos_hint).casefold()
        return "||".join([normalized_word, phonetic, word_type, pos_hint, meaning_hint, "phrase" if phrase else "word"])

    def _emit_duplicate_issues(
        self,
        duplicate_words: dict[str, list[WordRecord]],
        duplicate_phrases: dict[str, list[WordRecord]],
        source_type: str,
    ) -> list[ValidationIssue]:
        issues: list[ValidationIssue] = []
        for key, records in duplicate_words.items():
            if len(records) <= 1:
                continue
            references = self._format_duplicate_references(records)
            for record in records:
                issues.append(
                    self._issue(
                        category_key="word_duplicate_word",
                        record=record,
                        table_name="word",
                        field_name="word",
                        raw_value=record.word_text,
                        normalized_value=normalize_spaces(record.word_text).lower(),
                        reason=f"标准化后与其他单词重复，重复项: {references}。",
                        suggestion=(
                            "保留一条标准记录，其余记录删除或合并。"
                            if source_type == "db"
                            else "检查是否存在重复 Excel 行或重复导入。"
                        ),
                    )
                )

        for key, records in duplicate_phrases.items():
            if len(records) <= 1:
                continue
            references = self._format_duplicate_references(records)
            for record in records:
                issues.append(
                    self._issue(
                        category_key="word_duplicate_phrase",
                        record=record,
                        table_name="word",
                        field_name="word",
                        raw_value=record.word_text,
                        normalized_value=normalize_spaces(record.word_text).lower(),
                        reason=f"标准化后与其他词组重复，重复项: {references}。",
                        suggestion=(
                            "保留一条标准记录，其余记录删除或合并。"
                            if source_type == "db"
                            else "检查是否存在重复 Excel 行或重复导入。"
                        ),
                    )
                )
        return issues

    def _format_duplicate_references(self, records: list[WordRecord]) -> str:
        values: list[str] = []
        for record in records:
            if record.row_number is not None:
                values.append(f"{record.source_name}:第{record.row_number}行")
            elif record.record_id is not None:
                values.append(str(record.record_id))
        return ", ".join(values)

    def _issue(
        self,
        category_key: str,
        record: WordRecord,
        table_name: str,
        field_name: str,
        raw_value: str,
        normalized_value: str,
        reason: str,
        suggestion: str,
        auto_fix_value: str = "",
    ) -> ValidationIssue:
        return ValidationIssue(
            category_key=category_key,
            category_title=CATEGORY_TITLES[category_key],
            source_path=record.source_path,
            source_name=record.source_name,
            table_name=table_name,
            record_id=record.record_id,
            word_id=record.record_id,
            row_number=record.row_number,
            word_text=record.word_text,
            field_name=field_name,
            raw_value=to_text(raw_value),
            normalized_value=to_text(normalized_value),
            reason=reason,
            suggestion=suggestion,
            auto_fix_value=auto_fix_value,
        )

    def _issue_from_meaning(
        self,
        category_key: str,
        record: MeaningRecord,
        field_name: str,
        raw_value: str,
        normalized_value: str,
        reason: str,
        suggestion: str,
        auto_fix_value: str = "",
    ) -> ValidationIssue:
        return ValidationIssue(
            category_key=category_key,
            category_title=CATEGORY_TITLES[category_key],
            source_path=record.source_path,
            source_name=record.source_name,
            table_name="word_meaning",
            record_id=record.record_id,
            word_id=record.word_id,
            row_number=record.row_number,
            word_text=record.word_text,
            field_name=field_name,
            raw_value=to_text(raw_value),
            normalized_value=to_text(normalized_value),
            reason=reason,
            suggestion=suggestion,
            auto_fix_value=auto_fix_value,
        )

    def _build_fix_actions(self, issues: list[ValidationIssue]) -> list[FixAction]:
        safe_categories = {
            "word_word_spacing",
            "word_word_uppercase",
            "word_invalid_word_chars",
            "word_phrase_spacing",
            "word_phrase_uppercase",
            "word_invalid_phrase_chars",
            "phonetic_phrase_should_be_empty",
            "phonetic_not_standard_ipa",
            "meaning_zh_invalid_chars",
        }
        actions: dict[tuple[Path, int, str], FixAction] = {}
        for issue in issues:
            if issue.row_number is None:
                continue
            if issue.category_key not in safe_categories:
                continue
            if not issue.auto_fix_value and issue.category_key != "phonetic_phrase_should_be_empty":
                continue
            new_value = issue.auto_fix_value
            key = (issue.source_path, issue.row_number, issue.field_name)
            actions[key] = FixAction(
                source_path=issue.source_path,
                row_number=issue.row_number,
                field_name=issue.field_name,
                original_value=issue.raw_value,
                new_value=new_value,
                reason=issue.reason,
            )
        return sorted(actions.values(), key=lambda item: (str(item.source_path), item.row_number, item.field_name))


class ValidationWorker(QObject):
    finished = Signal(object)
    failed = Signal(str)
    log = Signal(str)

    def __init__(
        self,
        source_type: str,
        db_path: Path | None,
        excel_paths: list[Path],
        reference_db_path: Path | None,
        cache_path: Path,
        too_many_meanings: int,
    ) -> None:
        super().__init__()
        self.source_type = source_type
        self.db_path = db_path
        self.excel_paths = excel_paths
        self.reference_db_path = reference_db_path
        self.cache_path = cache_path
        self.too_many_meanings = too_many_meanings

    @Slot()
    def run(self) -> None:
        try:
            service = ValidationService(self.reference_db_path, self.cache_path, self.too_many_meanings)
            if self.source_type == "db":
                if self.db_path is None:
                    raise ValueError("数据库路径为空。")
                self.log.emit(f"开始校验数据库: {self.db_path}")
                summary = service.validate_database(self.db_path)
            else:
                self.log.emit(f"开始校验 Excel 文件，共 {len(self.excel_paths)} 个。")
                summary = service.validate_excel_files(self.excel_paths)
            for message in summary.log_messages:
                self.log.emit(message)
        except Exception as exc:
            self.failed.emit(str(exc))
            return
        self.finished.emit(summary)


class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("词库数据校验工具")
        self.resize(1520, 920)

        self.current_summary: ValidationSummary | None = None
        self.current_issues_by_category: dict[str, list[ValidationIssue]] = {}
        self.current_displayed_issues: list[ValidationIssue] = []
        self.worker_thread: QThread | None = None
        self.worker: ValidationWorker | None = None
        self.selected_excel_files: list[Path] = []
        self.confirmed_fix_actions: dict[tuple[Path, int, str], FixAction] = {}

        self.source_type_combo = QComboBox()
        self.source_type_combo.addItem("校验 words.db", "db")
        self.source_type_combo.addItem("校验 Excel 文件", "excel")

        self.db_path_edit = QLineEdit(str(DEFAULT_DB_PATH))
        self.db_browse_button = QPushButton("选择数据库")

        self.reference_db_edit = QLineEdit(str(DEFAULT_REFERENCE_DB_PATH))
        self.reference_db_button = QPushButton("参考词典")

        self.cache_path_edit = QLineEdit(str(DEFAULT_SEMANTIC_CACHE_PATH))

        self.too_many_meanings_spin = QSpinBox()
        self.too_many_meanings_spin.setRange(2, 100)
        self.too_many_meanings_spin.setValue(DEFAULT_TOO_MANY_MEANINGS)

        self.excel_select_button = QPushButton("选择 Excel 文件")
        self.excel_clear_button = QPushButton("清空列表")
        self.excel_list = QListWidget()
        self.excel_list.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)

        self.run_button = QPushButton("开始校验")
        self.confirm_fix_button = QPushButton("确认修复选中条目")
        self.confirm_fix_button.setEnabled(False)
        self.confirm_fix_button.setToolTip("仅 Excel 模式下可用，支持 Ctrl / Shift 多选后批量确认自动修复")
        self.fix_export_button = QPushButton("写回原 Excel")
        self.fix_export_button.setEnabled(False)
        self.fix_export_button.setToolTip("将已确认修复的条目直接写回原 Excel 文件，并自动生成备份")
        self.fix_status_label = QLabel("修复功能仅支持 Excel 模式。")
        self.fix_status_label.setWordWrap(True)
        self.summary_label = QLabel("等待开始")
        self.summary_label.setWordWrap(True)

        self.category_tree = QTreeWidget()
        self.category_tree.setColumnCount(2)
        self.category_tree.setHeaderLabels(["问题分类", "数量"])

        self.issue_table = QTableWidget(0, 10)
        self.issue_table.setHorizontalHeaderLabels(
            ["来源", "行/记录", "word_id", "单词/词组", "字段", "原始值", "标准值", "原因", "建议", "自动修复"]
        )
        self.issue_table.setAlternatingRowColors(True)
        self.issue_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.issue_table.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.issue_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.issue_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.issue_table.horizontalHeader().setStretchLastSection(True)
        self.issue_table.verticalHeader().setVisible(False)

        self.log_output = QPlainTextEdit()
        self.log_output.setReadOnly(True)
        self.log_output.setPlaceholderText("运行日志会显示在这里")

        self._build_layout()
        self._connect_signals()
        self._update_source_visibility()

    def _build_layout(self) -> None:
        controls_layout = QGridLayout()
        controls_layout.addWidget(QLabel("数据源类型"), 0, 0)
        controls_layout.addWidget(self.source_type_combo, 0, 1)
        controls_layout.addWidget(QLabel("参考词典 DB"), 0, 2)
        controls_layout.addWidget(self.reference_db_edit, 0, 3)
        controls_layout.addWidget(self.reference_db_button, 0, 4)

        controls_layout.addWidget(QLabel("数据库路径"), 1, 0)
        controls_layout.addWidget(self.db_path_edit, 1, 1, 1, 3)
        controls_layout.addWidget(self.db_browse_button, 1, 4)

        controls_layout.addWidget(QLabel("语义缓存 JSON"), 2, 0)
        controls_layout.addWidget(self.cache_path_edit, 2, 1, 1, 3)
        controls_layout.addWidget(QLabel("多义词阈值"), 2, 4)
        controls_layout.addWidget(self.too_many_meanings_spin, 2, 5)

        excel_tools_layout = QHBoxLayout()
        excel_tools_layout.addWidget(self.excel_select_button)
        excel_tools_layout.addWidget(self.excel_clear_button)
        excel_tools_layout.addWidget(QLabel("支持一次选择多个 .xls / .xlsx 文件"), 1)

        action_layout = QHBoxLayout()
        action_layout.addWidget(self.run_button)
        action_layout.addWidget(self.confirm_fix_button)
        action_layout.addWidget(self.fix_export_button)
        action_layout.addStretch(1)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.addWidget(self.category_tree)
        splitter.addWidget(self.issue_table)
        splitter.setStretchFactor(0, 2)
        splitter.setStretchFactor(1, 7)

        central_widget = QWidget()
        root_layout = QVBoxLayout(central_widget)
        root_layout.addLayout(controls_layout)
        root_layout.addLayout(excel_tools_layout)
        root_layout.addWidget(self.excel_list, 1)
        root_layout.addLayout(action_layout)
        root_layout.addWidget(self.fix_status_label)
        root_layout.addWidget(QLabel("校验结果"))
        root_layout.addWidget(self.summary_label)
        root_layout.addWidget(splitter, 5)
        root_layout.addWidget(QLabel("运行日志"))
        root_layout.addWidget(self.log_output, 2)
        self.setCentralWidget(central_widget)

    def _connect_signals(self) -> None:
        self.source_type_combo.currentIndexChanged.connect(self._update_source_visibility)
        self.db_browse_button.clicked.connect(self.select_db_path)
        self.reference_db_button.clicked.connect(self.select_reference_db)
        self.excel_select_button.clicked.connect(self.select_excel_files)
        self.excel_clear_button.clicked.connect(self.clear_excel_files)
        self.run_button.clicked.connect(self.start_validation)
        self.confirm_fix_button.clicked.connect(self.confirm_current_issue_fix)
        self.fix_export_button.clicked.connect(self.export_fixed_excel_files)
        self.category_tree.currentItemChanged.connect(self.on_category_changed)
        self.issue_table.itemSelectionChanged.connect(self.update_fix_button_state)

    def _update_source_visibility(self) -> None:
        source_type = self.current_source_type
        is_db = source_type == "db"
        self.db_path_edit.setEnabled(is_db)
        self.db_browse_button.setEnabled(is_db)
        self.excel_select_button.setEnabled(not is_db)
        self.excel_clear_button.setEnabled(not is_db)
        self.excel_list.setEnabled(not is_db)
        self.update_fix_button_state()

    @property
    def current_source_type(self) -> str:
        return str(self.source_type_combo.currentData())

    @Slot()
    def select_db_path(self) -> None:
        selected, _ = QFileDialog.getOpenFileName(
            self,
            "选择 words.db",
            str(Path(self.db_path_edit.text()).parent),
            "SQLite Database (*.db *.sqlite *.sqlite3);;All Files (*)",
        )
        if selected:
            self.db_path_edit.setText(selected)

    @Slot()
    def select_reference_db(self) -> None:
        selected, _ = QFileDialog.getOpenFileName(
            self,
            "选择参考词典数据库",
            str(Path(self.reference_db_edit.text()).parent),
            "SQLite Database (*.db *.sqlite *.sqlite3);;All Files (*)",
        )
        if selected:
            self.reference_db_edit.setText(selected)

    @Slot()
    def select_excel_files(self) -> None:
        files, _ = QFileDialog.getOpenFileNames(
            self,
            "选择一个或多个 Excel 文件",
            str(DEFAULT_DB_PATH.parent / "excel"),
            "Excel 文件 (*.xlsx *.xls);;所有文件 (*.*)",
        )
        if not files:
            return
        self.selected_excel_files = [Path(file) for file in files if not Path(file).name.startswith("~$")]
        self.excel_list.clear()
        for file in self.selected_excel_files:
            self.excel_list.addItem(str(file))
        self.append_log(f"已选择 {len(self.selected_excel_files)} 个 Excel 文件。")

    @Slot()
    def clear_excel_files(self) -> None:
        self.selected_excel_files = []
        self.excel_list.clear()
        self.append_log("已清空 Excel 文件列表。")

    @Slot()
    def start_validation(self) -> None:
        if self.worker_thread is not None:
            QMessageBox.information(self, "提示", "校验任务正在运行，请等待完成。")
            return

        source_type = self.current_source_type
        db_path = Path(self.db_path_edit.text().strip()) if self.db_path_edit.text().strip() else None
        reference_db = Path(self.reference_db_edit.text().strip()) if self.reference_db_edit.text().strip() else None
        cache_path = Path(self.cache_path_edit.text().strip()) if self.cache_path_edit.text().strip() else DEFAULT_SEMANTIC_CACHE_PATH

        if source_type == "db":
            if db_path is None or not db_path.exists():
                QMessageBox.warning(self, "路径错误", f"数据库不存在:\n{db_path}")
                return
        else:
            if not self.selected_excel_files:
                QMessageBox.warning(self, "提示", "请先选择 Excel 文件。")
                return

        self.current_summary = None
        self.current_issues_by_category = {}
        self.current_displayed_issues = []
        self.confirmed_fix_actions = {}
        self.category_tree.clear()
        self.issue_table.setRowCount(0)
        self.summary_label.setText("校验中，请稍候...")
        self.run_button.setEnabled(False)
        self.confirm_fix_button.setEnabled(False)
        self.fix_export_button.setEnabled(False)
        self.log_output.clear()

        self.worker_thread = QThread(self)
        self.worker = ValidationWorker(
            source_type=source_type,
            db_path=db_path,
            excel_paths=list(self.selected_excel_files),
            reference_db_path=reference_db if reference_db and reference_db.exists() else None,
            cache_path=cache_path,
            too_many_meanings=int(self.too_many_meanings_spin.value()),
        )
        self.worker.moveToThread(self.worker_thread)
        self.worker_thread.started.connect(self.worker.run)
        self.worker.finished.connect(self.on_validation_finished)
        self.worker.failed.connect(self.on_validation_failed)
        self.worker.log.connect(self.append_log)
        self.worker.finished.connect(self.worker_thread.quit)
        self.worker.failed.connect(self.worker_thread.quit)
        self.worker_thread.finished.connect(self.cleanup_worker)
        self.worker_thread.start()

    @Slot(object)
    def on_validation_finished(self, summary: ValidationSummary) -> None:
        self.current_summary = summary
        self.current_displayed_issues = []
        self.confirmed_fix_actions = {}
        issues_by_category: dict[str, list[ValidationIssue]] = defaultdict(list)
        for issue in summary.issues:
            issues_by_category[issue.category_key].append(issue)
        self.current_issues_by_category = dict(
            sorted(issues_by_category.items(), key=lambda item: (-len(item[1]), item[0]))
        )

        self.category_tree.clear()
        root_item = QTreeWidgetItem(["全部问题", str(len(summary.issues))])
        root_item.setData(0, Qt.ItemDataRole.UserRole, "__all__")
        self.category_tree.addTopLevelItem(root_item)

        for category_key, issues in self.current_issues_by_category.items():
            item = QTreeWidgetItem([CATEGORY_TITLES.get(category_key, category_key), str(len(issues))])
            item.setData(0, Qt.ItemDataRole.UserRole, category_key)
            root_item.addChild(item)

        root_item.setExpanded(True)
        self.category_tree.setCurrentItem(root_item)

        if not summary.issues:
            self.summary_label.setText(
                f"校验完成。来源: {summary.source_label}，word {summary.total_words} 条，meaning {summary.total_meanings} 条，未发现问题。"
            )
        else:
            self.summary_label.setText(
                f"校验完成。来源: {summary.source_label}，word {summary.total_words} 条，meaning {summary.total_meanings} 条，"
                f"共发现 {len(summary.issues)} 条问题，分布在 {len(self.current_issues_by_category)} 类原因中。"
            )

        self.confirm_fix_button.setEnabled(False)
        self.fix_export_button.setEnabled(False)

    @Slot(str)
    def on_validation_failed(self, message: str) -> None:
        self.summary_label.setText("校验失败")
        self.append_log(f"校验失败: {message}")
        QMessageBox.critical(self, "校验失败", message)

    @Slot()
    def cleanup_worker(self) -> None:
        if self.worker is not None:
            self.worker.deleteLater()
        if self.worker_thread is not None:
            self.worker_thread.deleteLater()
        self.worker = None
        self.worker_thread = None
        self.run_button.setEnabled(True)
        self.update_fix_button_state()

    @Slot(str)
    def append_log(self, message: str) -> None:
        self.log_output.appendPlainText(message)

    @Slot(QTreeWidgetItem, QTreeWidgetItem)
    def on_category_changed(self, current: QTreeWidgetItem | None, _: QTreeWidgetItem | None) -> None:
        if current is None or self.current_summary is None:
            self.populate_issue_table([])
            return

        category_key = current.data(0, Qt.ItemDataRole.UserRole)
        if category_key == "__all__":
            self.populate_issue_table(self.current_summary.issues)
            return
        self.populate_issue_table(self.current_issues_by_category.get(str(category_key), []))

    def populate_issue_table(self, issues: list[ValidationIssue]) -> None:
        self.current_displayed_issues = issues
        self.issue_table.setRowCount(len(issues))
        for row_index, issue in enumerate(issues):
            row_or_id = str(issue.row_number) if issue.row_number is not None else ("" if issue.record_id is None else str(issue.record_id))
            values = [
                issue.source_name,
                row_or_id,
                "" if issue.word_id is None else str(issue.word_id),
                issue.word_text,
                issue.field_name,
                issue.raw_value,
                issue.normalized_value,
                issue.reason,
                issue.suggestion,
                issue.auto_fix_value,
            ]
            for column_index, value in enumerate(values):
                item = QTableWidgetItem(value)
                if column_index in {1, 2}:
                    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.issue_table.setItem(row_index, column_index, item)

            color = self._issue_color(issue)
            for column_index in range(self.issue_table.columnCount()):
                cell = self.issue_table.item(row_index, column_index)
                if cell is not None:
                    cell.setBackground(color)
        self.issue_table.resizeColumnsToContents()
        self.update_fix_button_state()

    def _issue_color(self, issue: ValidationIssue) -> QColor:
        fix_key = self._issue_fix_key(issue)
        if fix_key is not None and fix_key in self.confirmed_fix_actions:
            return QColor("#DFF5E1")
        return COLOR_BY_PREFIX.get(issue.category_key.split("_", 1)[0], QColor("#F7F7F7"))

    def _issue_fix_key(self, issue: ValidationIssue) -> tuple[Path, int, str] | None:
        if issue.row_number is None:
            return None
        return (issue.source_path, issue.row_number, issue.field_name)

    def _build_fix_action_from_issue(self, issue: ValidationIssue) -> FixAction | None:
        if issue.row_number is None:
            return None
        if self.current_summary is None or self.current_summary.source_type != "excel":
            return None
        safe_categories = {
            "word_word_spacing",
            "word_word_uppercase",
            "word_invalid_word_chars",
            "word_phrase_spacing",
            "word_phrase_uppercase",
            "word_invalid_phrase_chars",
            "phonetic_phrase_should_be_empty",
            "phonetic_not_standard_ipa",
            "meaning_zh_invalid_chars",
        }
        if issue.category_key not in safe_categories:
            return None
        if not issue.auto_fix_value and issue.category_key != "phonetic_phrase_should_be_empty":
            return None
        return FixAction(
            source_path=issue.source_path,
            row_number=issue.row_number,
            field_name=issue.field_name,
            original_value=issue.raw_value,
            new_value=issue.auto_fix_value,
            reason=issue.reason,
        )

    @Slot()
    def update_fix_button_state(self) -> None:
        selected_issues = self.get_selected_issues()
        issue = selected_issues[0] if selected_issues else None
        fixable_count = sum(1 for selected_issue in selected_issues if self._build_fix_action_from_issue(selected_issue) is not None)
        self.confirm_fix_button.setEnabled(fixable_count > 0)
        self.fix_export_button.setEnabled(
            bool(self.confirmed_fix_actions) and self.current_summary is not None and self.current_summary.source_type == "excel"
        )
        self.fix_status_label.setText(self._build_fix_status_text(issue, fixable_count, len(selected_issues)))

    def _build_fix_status_text(self, issue: ValidationIssue | None, fixable_count: int, selected_count: int) -> str:
        if self.current_source_type != "excel":
            return "当前是数据库模式。修复功能只对 Excel 校验结果开放。"
        if self.current_summary is None:
            return "请选择 Excel 文件并执行校验后，再确认需要修复的条目。"
        if not self.current_summary.issues:
            return "当前没有问题，不需要修复。"
        if issue is None:
            if self.confirmed_fix_actions:
                return f"已确认 {len(self.confirmed_fix_actions)} 条修复，可继续选择条目或直接写回原 Excel。"
            return "请先在右侧问题表中选中支持自动修复的问题，支持 Shift 连选或 Ctrl 多选。"
        if fixable_count > 0:
            if selected_count == 1:
                return "当前选中条目支持自动修复，可点击“确认修复选中条目”。"
            if fixable_count == selected_count:
                return f"当前已选中 {selected_count} 条问题，均支持自动修复，可批量确认。"
            return f"当前已选中 {selected_count} 条问题，其中 {fixable_count} 条支持自动修复，确认时将跳过其余条目。"
        if self.confirmed_fix_actions:
            return f"当前选中条目不支持自动修复。已确认 {len(self.confirmed_fix_actions)} 条修复，可直接写回原 Excel。"
        return "当前选中条目不支持自动修复。只有空格、大小写、非法字符、词组音标清空、音标标准化、释义非法字符清理等问题可自动修复。"

    def get_selected_issues(self) -> list[ValidationIssue]:
        selection_model = self.issue_table.selectionModel()
        if selection_model is None:
            return []

        issues: list[ValidationIssue] = []
        for index in sorted(selection_model.selectedRows(), key=lambda selected_index: selected_index.row()):
            row = index.row()
            if 0 <= row < len(self.current_displayed_issues):
                issues.append(self.current_displayed_issues[row])
        return issues

    def get_selected_issue(self) -> ValidationIssue | None:
        selected_issues = self.get_selected_issues()
        if not selected_issues:
            return None
        return selected_issues[0]

    @Slot()
    def confirm_current_issue_fix(self) -> None:
        selected_issues = self.get_selected_issues()
        if not selected_issues:
            QMessageBox.information(self, "提示", "请先在表格中选择至少一条问题记录。")
            return

        new_actions: dict[tuple[Path, int, str], FixAction] = {}
        skipped_count = 0
        for issue in selected_issues:
            action = self._build_fix_action_from_issue(issue)
            if action is None:
                skipped_count += 1
                continue
            fix_key = (action.source_path, action.row_number, action.field_name)
            new_actions[fix_key] = action

        if not new_actions:
            QMessageBox.information(self, "提示", "当前选中的问题都不支持自动修复。")
            return

        added_count = 0
        for fix_key, action in new_actions.items():
            if fix_key not in self.confirmed_fix_actions:
                added_count += 1
            self.confirmed_fix_actions[fix_key] = action
            self.append_log(
                f"已确认修复: {action.source_path.name} 第 {action.row_number} 行 {action.field_name} -> {action.new_value}"
            )

        total_fixable = len(new_actions)
        if len(selected_issues) == 1 and skipped_count == 0:
            QMessageBox.information(self, "提示", "当前条目已加入待写回修复列表。")
        else:
            message = f"已处理 {len(selected_issues)} 条选中问题，其中 {total_fixable} 条支持自动修复"
            if added_count != total_fixable:
                message += f"，新增 {added_count} 条"
            if skipped_count > 0:
                message += f"，跳过 {skipped_count} 条不支持自动修复的问题"
            message += "。"
            QMessageBox.information(self, "批量确认完成", message)

        self.populate_issue_table(self.current_displayed_issues)

    @Slot()
    def export_fixed_excel_files(self) -> None:
        if self.current_summary is None or self.current_summary.source_type != "excel":
            QMessageBox.information(self, "提示", "当前结果没有可写回的 Excel 自动修复项。")
            return
        if not self.confirmed_fix_actions:
            QMessageBox.information(self, "提示", "请先确认至少一条需要修复的问题，再写回原 Excel。")
            return

        reply = QMessageBox.question(
            self,
            "确认写回",
            "将直接修改原 Excel 文件，并在同目录生成 .bak 备份文件。\n是否继续？",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return

        reference_db = Path(self.reference_db_edit.text().strip()) if self.reference_db_edit.text().strip() else None
        service = ValidationService(
            reference_db_path=reference_db if reference_db and reference_db.exists() else None,
            cache_path=Path(self.cache_path_edit.text().strip()) if self.cache_path_edit.text().strip() else DEFAULT_SEMANTIC_CACHE_PATH,
            too_many_meanings=int(self.too_many_meanings_spin.value()),
        )
        try:
            updated_files = service.apply_fixed_excel_files(list(self.confirmed_fix_actions.values()))
        except Exception as exc:
            QMessageBox.critical(self, "写回失败", str(exc))
            return

        if not updated_files:
            QMessageBox.information(self, "提示", "没有成功写回的 Excel 文件。")
            return

        self.append_log(f"已写回 {len(updated_files)} 个 Excel 文件，并生成备份。")
        QMessageBox.information(
            self,
            "写回完成",
            "以下 Excel 已写回原文件，并生成备份:\n" + "\n".join(str(path) for path in updated_files),
        )


def main() -> int:
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())