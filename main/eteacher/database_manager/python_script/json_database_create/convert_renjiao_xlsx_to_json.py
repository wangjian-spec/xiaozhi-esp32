from __future__ import annotations

import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
BOOK_DIR = BASE_DIR / "book"
SOURCE_XLSX_PATH = BOOK_DIR / "人教版小学单词表（3-6年级）.xlsx"
FALLBACK_SCHEMA_PATHS = [
    BASE_DIR / "words_stage1.json",
    BOOK_DIR / "words_stage1.json",
    BOOK_DIR / "沪教牛津版小学单词表（3-6年级）.json",
    BOOK_DIR / "外研版小学单词表（1-6年级）.json",
]
OUTPUT_JSON_PATH = BOOK_DIR / "人教版小学单词表（3-6年级）.json"
PUBLISHER = "人教版"
EXPECTED_HEADERS = ["课本", "主题", "单词", "音标", "中文"]
WHITESPACE_PATTERN = re.compile(r"\s+")


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    return WHITESPACE_PATTERN.sub(" ", text)


def normalize_phonetic(value: Any) -> str | None:
    phonetic = normalize_text(value)
    if not phonetic:
        return None
    return phonetic.replace("'", "ˈ")


def load_rows(xlsx_path: Path) -> list[list[Any]]:
    workbook = load_workbook(xlsx_path, data_only=True)
    sheet = workbook.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def detect_header_map(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(rows[:5]):
        normalized = [normalize_text(cell) for cell in row]
        mapping = {
            header: normalized.index(header)
            for header in EXPECTED_HEADERS
            if header in normalized
        }
        if len(mapping) == len(EXPECTED_HEADERS):
            return row_index, mapping
    raise ValueError(f"前 5 行内未找到完整表头，要求字段: {', '.join(EXPECTED_HEADERS)}")


def build_source(book_label: str, topic: str) -> str:
    parts = [PUBLISHER]
    if book_label:
        parts.append(book_label)
    if topic:
        parts.append(topic)
    return " | ".join(parts)


def load_reference_schema() -> dict[str, Any]:
    errors: list[str] = []
    for path in FALLBACK_SCHEMA_PATHS:
        if not path.exists():
            errors.append(f"{path}: 文件不存在")
            continue

        payload: dict[str, Any] | None = None
        for encoding in ("utf-8", "utf-8-sig"):
            try:
                payload = json.loads(path.read_text(encoding=encoding))
                break
            except (UnicodeDecodeError, json.JSONDecodeError) as exc:
                errors.append(f"{path} ({encoding}): {exc}")

        if not isinstance(payload, dict):
            continue

        schema = payload.get("schema")
        if isinstance(schema, dict) and all(
            table_name in schema
            for table_name in ("word", "word_meaning", "word_form", "word_example")
        ):
            return schema

        errors.append(f"{path}: 缺少完整 schema")

    raise ValueError("无法从参考 JSON 加载 schema:\n" + "\n".join(errors))


def parse_records(rows: list[list[Any]]) -> tuple[list[dict[str, Any]], int]:
    header_row_index, mapping = detect_header_map(rows)
    records: list[dict[str, Any]] = []
    word_id_by_key: dict[str, int] = {}
    word_payload_by_key: dict[str, dict[str, Any]] = {}
    next_word_id = 1
    next_meaning_id = 1

    for row_number, row in enumerate(rows[header_row_index + 1 :], start=header_row_index + 2):
        if all(not normalize_text(cell) for cell in row):
            continue

        book_label = normalize_text(row[mapping["课本"]])
        topic = normalize_text(row[mapping["主题"]])
        word = normalize_text(row[mapping["单词"]])
        phonetic = normalize_phonetic(row[mapping["音标"]])
        meaning_zh = normalize_text(row[mapping["中文"]])

        if not word:
            continue
        if not meaning_zh:
            raise ValueError(f"第 {row_number} 行缺少中文释义: {row}")

        word_key = word.casefold()
        if word_key not in word_id_by_key:
            word_id_by_key[word_key] = next_word_id
            word_payload_by_key[word_key] = {
                "id": next_word_id,
                "word": word,
                "phonetic": phonetic,
                "word_type": None,
            }
            next_word_id += 1
        else:
            existing_word = word_payload_by_key[word_key]
            if not existing_word.get("phonetic") and phonetic:
                existing_word["phonetic"] = phonetic

        word_id = word_id_by_key[word_key]
        meaning_id = next_meaning_id
        next_meaning_id += 1

        records.append(
            {
                "word": dict(word_payload_by_key[word_key]),
                "word_meaning": {
                    "id": meaning_id,
                    "word_id": word_id,
                    "stage": 1,
                    "pos": None,
                    "meaning_en": None,
                    "meaning_zh": meaning_zh,
                    "source": build_source(book_label, topic),
                    "word_tag": None,
                },
                "word_form": [],
                "word_example": [],
            }
        )
        records[-1]["word"].setdefault("image", None)

    return records, len(word_id_by_key)


def build_output(records: list[dict[str, Any]], unique_word_count: int) -> dict[str, Any]:
    schema = load_reference_schema()
    return {
        "metadata": {
            "source_file": str(SOURCE_XLSX_PATH),
            "source_workbook": str(SOURCE_XLSX_PATH),
            "exported_at": datetime.now().isoformat(timespec="seconds"),
            "record_count": len(records),
            "root_table": "word_meaning",
            "related_tables": ["word", "word_meaning", "word_form", "word_example"],
            "unique_word_count": unique_word_count,
            "publisher": PUBLISHER,
        },
        "schema": schema,
        "records": records,
    }


def main() -> int:
    if not SOURCE_XLSX_PATH.exists():
        raise FileNotFoundError(f"找不到源文件: {SOURCE_XLSX_PATH}")

    rows = load_rows(SOURCE_XLSX_PATH)
    records, unique_word_count = parse_records(rows)
    output = build_output(records, unique_word_count)
    OUTPUT_JSON_PATH.write_text(
        json.dumps(output, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(
        f"已生成 {OUTPUT_JSON_PATH}，共 {len(records)} 条 records，"
        f"去重后 {unique_word_count} 个唯一单词。"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())